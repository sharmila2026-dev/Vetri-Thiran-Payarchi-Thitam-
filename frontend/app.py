import io
import re
import os
import requests
import streamlit as st
from docx import Document
from docx.shared import Inches, Pt, RGBColor
from fpdf import FPDF
from config import LOGO_PATH, INVERSE_LOGO_PATH

st.set_page_config(page_title="LegalEase", layout="centered")

def sanitize_text(text: str) -> str:
    text = text.replace("“", '"').replace("”", '"').replace("’", "'").replace("‘", "'")
    return text

def format_docx(text: str, doc_type: str) -> io.BytesIO:
    doc = Document()
    
    if os.path.exists(LOGO_PATH):
        doc.add_image(LOGO_PATH, width=Inches(1.5))
    
    title_p = doc.add_paragraph()
    run = title_p.add_run(f"\n{doc_type}\n")
    run.bold = True
    run.font.size = Pt(18)
    
    for paragraph in text.split("\n\n"):
        if paragraph.startswith("## ") or paragraph.startswith("# "):
            p = doc.add_paragraph()
            r = p.add_run(paragraph.replace("#", "").strip())
            r.bold = True
            r.font.size = Pt(14)
        else:
            p = doc.add_paragraph(paragraph.strip())
            p.style.font.name = 'Times New Roman'
            p.style.font.size = Pt(11)
            
    section = doc.sections[0]
    footer = section.footer
    f_p = footer.paragraphs[0]
    f_p.text = "LegalEase Inc. | contact@legalease.com | All Rights Reserved."
    f_p.style.font.size = Pt(9)
    
    buffer = io.BytesIO()
    doc.save(buffer)
    buffer.seek(0)
    return buffer

class LegalPDF(FPDF):
    def header(self):
        if os.path.exists(LOGO_PATH):
            self.image(LOGO_PATH, 10, 8, 33)
        self.set_font('Arial', 'B', 12)
        self.cell(0, 10, 'LegalEase Document', 0, 1, 'R')
        self.ln(5)

    def footer(self):
        self.set_y(-15)
        self.set_font('Arial', 'I', 8)
        self.cell(0, 10, 'LegalEase Inc. | contact@legalease.com | All Rights Reserved.', 0, 0, 'C')

def format_pdf(text: str, doc_type: str) -> bytes:
    pdf = LegalPDF()
    pdf.add_page()
    pdf.set_auto_page_break(auto=True, margin=15)
    
    pdf.set_font("Arial", 'B', 16)
    pdf.cell(0, 10, doc_type, 0, 1, 'C')
    pdf.ln(5)
    
    clean_text = sanitize_text(text)
    
    pdf.set_font("Arial", size=10)
    for line in clean_text.split('\n'):
        if line.startswith("#"):
            pdf.set_font("Arial", 'B', 12)
            pdf.multi_cell(0, 8, line.replace("#", "").strip())
            pdf.set_font("Arial", size=10)
        else:
            pdf.multi_cell(0, 6, line)
        pdf.ln(2)
        
    return pdf.output(dest='S').encode('latin-1', errors='replace')

def format_html_preview(text: str) -> str:
    formatted_html = text.replace('\n', '